from django.shortcuts import render, redirect
from django.http import JsonResponse
from django.utils.safestring import mark_safe
from django.contrib.auth.forms import UserCreationForm
from .models import Session, Profile, Question, QuestionAttempt, PageView, WebsiteSession, User, CursorEvent, DeletionEvent
import json
from django.views.decorators.csrf import csrf_exempt
from .chatbot import generateResponse
from .forms import StyledUserCreationForm, StyledAuthenticationForm
from django.contrib.auth import login, logout
from .compiler import compileCode, handler
from .question_assignment import initialQuestion, assignment
from django.views.decorators.http import require_POST
from django.http import JsonResponse, HttpRequest
from json import JSONDecodeError
from django.db import transaction
from django.db.models import F
from django.utils import timezone


def home(request): 
    return render(request, "studentSystem/index.html")

def signupView(request): # TODO - Style error messages
    if request.method == "POST":
        form = StyledUserCreationForm(request.POST)
        if form.is_valid():
            new_user = form.save()
            login(request, new_user)

            # Call initialization logic
            currentProfile = Profile.objects.get(user=new_user)
            initialQuestion(currentProfile)

            # Queue the first set of questions
            nextQuestions = currentProfile.next_questions.all()

            if nextQuestions.exists():
                queuedQuestion = nextQuestions.first()
                session = Session.objects.create(user=new_user, question=queuedQuestion)

                session.codeSubmission = (
                    queuedQuestion.target_function_definition +
                    " (* <- Keep the provided function name. It needs to be maintained for the compiler to work properly. *)"
                )
                session.save()

                attempt = QuestionAttempt.objects.create(
                    associated_profile=currentProfile,
                    question=queuedQuestion,
                    solved_test_cases=0,
                    session=session
                )
                currentProfile.question_history.add(attempt)
                currentProfile.next_questions.remove(queuedQuestion)

            # Render the dashboard like a GET
            return redirect("/dashboard/")  

        else:
            pass
    else:
        form = StyledUserCreationForm()
    return render(request, "studentSystem/signup.html", {"form": form})


def loginView(request): # TODO - Style error messages
    if request.method == "POST":
        form = StyledAuthenticationForm(data=request.POST)
        if form.is_valid():
            loggedUser = form.get_user()
            login(request, loggedUser)

            session = Session.objects.filter(user=loggedUser).first()
            return redirect("dashboard")
        else:
            pass
    else:
        form = StyledAuthenticationForm()
    return render(request, "studentSystem/login.html", { "form" : form})

def logoutView(request):
    logout(request) 
    return redirect("home") 
 
def dashboard(request): # TODO: Style "# of cases passed" tooltips to float above past problem blocks
    currentUser = request.user
    currentProfile = Profile.objects.get(user=request.user) # Get the current user's profile
    sessions = Session.objects.filter(user=currentUser).order_by('-pk') # All sessions associated with the user, ordered by most recent
    nextQuestions = currentProfile.next_questions.all()

    request_type = None
    if request.method == "POST": # Use POST request to request the next question
        if request.content_type == "application/json":
            body = json.loads(request.body.decode("utf-8"))
            request_type = body.get("request_type")
        else:
            request_type = request.POST.get("request_type")

        if request_type == "next_question" and nextQuestions.exists():
            queuedQuestion = nextQuestions.first()

            attempt = QuestionAttempt.objects.create(
                associated_profile=currentProfile,
                question=queuedQuestion,
                solved_test_cases=0,
                session=Session.objects.create(user=currentUser, question=queuedQuestion)
            )

            currentProfile.question_history.add(attempt)
            currentProfile.next_questions.remove(nextQuestions.first())

        elif request_type == "next_question" and not nextQuestions.exists():
            assignment(currentProfile)
            queuedQuestion = nextQuestions.first()

            attempt = QuestionAttempt.objects.create(
                associated_profile=currentProfile,
                question=queuedQuestion,
                solved_test_cases=0,
                session=Session.objects.create(user=currentUser, question=queuedQuestion)
            )

            currentProfile.question_history.add(attempt)
            currentProfile.next_questions.remove(nextQuestions.first())

        currentQuestionAttempt = currentProfile.question_history.latest('created_at')
        currentQuestionAttempt.session.codeSubmission = currentQuestionAttempt.question.target_function_definition + " (* <- Keep the provided function name. It needs to be maintained for the compiler to work properly. *)"
        currentQuestionAttempt.session.save()

    currentQuestionAttempt = currentProfile.question_history.latest('created_at')
    
    pastQuestions = currentProfile.question_history.all().order_by('-pk') [1:] # Exclude the current question from past questions & order by most recent

    context = {
        "sessions" : sessions,
        "user" : currentUser,
        "pastQuestions" : pastQuestions,
        "currentQuestion" : currentQuestionAttempt,
    }
    return render(request, "studentSystem/dashboard.html", context)

def createChat (request):
    currentUser = request.user
    session = Session.objects.create(user=currentUser)
    return redirect('chat', id=session.pk)

def chat(request, id):
    session = Session.objects.get(pk = id) # Current session and associated fields
    chatLog = json.dumps(session.chatLog) # Chats of the current session
    existingChats = Session.objects.filter(user=session.user).exclude(pk=id) # Other chats asssociated with the user

    context = {
        "session" : session,
        "sessionID" : id,
        "chatLog" : chatLog,
        "existingChats" : existingChats
    }
    return render(request, "studentSystem/chat.html", context)

def updateChat(request, id):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            new_chat = data.get('chat')
            session = Session.objects.get(pk = id)

            if not isinstance(session.chatLog, list):
                session.chatLog = []
            session.chatLog.append(new_chat) 
            session.save()
            return JsonResponse({'status': 'success'})
        except Session.DoesNotExist:
            return JsonResponse({'status': 'error', 'message': 'Item not found'}, status=404)
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
    else:
        return JsonResponse({'status': 'error', 'message': 'Invalid request'}, status=400)
    
def getAIResponse(request, id):
    if request.method == "POST":
        session = Session.objects.get(pk = id)
        conversationHistory = session.chatLog
        conversationHistory.insert(0, {"role": "system", "content": "You are a helpful programming assistant that has expertise in StandardML."})
        conversationHistory.insert(1, {"role": "assistant", "content": "Hello, I am a helpful programming assistant specialized in StandardML! Let me know if you have any questions about the language, and I'm here to help!"})
        for entry in range(2, len(conversationHistory)):
            if entry % 2 == 1:
                conversationHistory[entry] = {"role": "assistant", "content": conversationHistory[entry]}
            else:
                conversationHistory[entry] = {"role": "user", "content": conversationHistory[entry]}

        body = json.loads(request.body)
        prompt = (body.get('prompt')) 
        conversationHistory.append({"role": "user", "content": prompt})

        ai_response = generateResponse(conversationHistory)
        data = {"message" : ai_response}
        return JsonResponse(data)
    
def codeSubmit(request, id): # TODO: Add tooltips to buttons
    session = Session.objects.get(pk = id) 
    submission = session.codeSubmission
    context = {
        "codeSubmission" : submission,
        "session" : session,
        "nextSession" : id + 1,
        "previousSession" : id - 1,
    }
    return render(request, "studentSystem/codeSubmit.html", context)

@csrf_exempt
def saveSubmission(request, id):
    try:
        session = Session.objects.get(pk=id)
    except Session.DoesNotExist:
        return JsonResponse({'status': 'error', 'message': 'Item not found'}, status=404)

    if request.method == "POST":
        try:
            body = json.loads(request.body)
            new_submission = body.get("submission")

            session.codeSubmission = new_submission
            session.save()

            return JsonResponse({'status': 'success'})
        except Exception as e:
            print("Parse error:", e)
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)

    return JsonResponse({'status': 'error', 'message': 'Invalid request'}, status=400)

def askAI(request, id):
    if request.method == "POST":

        body = json.loads(request.body)

        submission = body.get("submission")
        compiler_output = body.get("compiler_output")

        response = generateResponse([{"role" : "system", "content" : "You are a helpful programming assistant that has expertise in StandardML."},
                                      {"role" : "user", "content" : "Analyze the following StandardML code: " + submission
                                    +"This what the code currently outputs in the compiler when it is run: " + compiler_output
                                    +"If there is an error, explain what the error means and how to fix it." + 
                                    " If there is no error, suggest improvements to the code or optimizations that can be made." }])
        return JsonResponse({"AI_response" : response})

def runCompile(request): 
    if request.method == "POST":
        body = json.loads(request.body)
        submission = body.get("submission")
        compiled_output = str(compileCode(submission))
        return JsonResponse({"compiled_output" : compiled_output})

def submitCode(request, id):
    if request.method == "POST": 
        currentProfile = Profile.objects.get(user=request.user)
        session = Session.objects.get(pk=id)

        casesSolved = handler(currentProfile, session=session)

        return JsonResponse({"status": "ok", "casesSolved": casesSolved})

    return JsonResponse({"status": "error", "message": "Invalid request"}, status=400)

def user_survey(request, id):
    session = Session.objects.get(pk = id)
    attempt = QuestionAttempt.objects.get(session=session)
    solved_test_cases = attempt.solved_test_cases
    context = {
        "session_id" : id,
        "solved_test_cases": solved_test_cases
    }
    return render(request, "studentSystem/userSurvey.html", context)

def post_survey_response(request, id):
    if request.method == "POST":
        print ("Successfully got to view")
        body = json.loads(request.body)
        survey_response = body.get("survey_response")
        session = Session.objects.get(pk = id)
        attempt = QuestionAttempt.objects.get(session=session)
        attempt.survey_response = survey_response
        attempt.save()
        return JsonResponse({"status": "ok"})
    else:
        return JsonResponse({"status": "error", "message": "Invalid request"}, status=400)


@csrf_exempt  # Using sendBeacon, so CSRF token is not available
def track_page_time(request):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "POST only"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
    except Exception:
        return JsonResponse({"ok": False, "error": "invalid json"}, status=400)

    # Total time from page load until submit, in ms
    duration_ms = int(max(0, data.get("duration_ms", 0)))
    path = data.get("path") or request.path

    started_at_ms = data.get("started_at_ms")
    if started_at_ms:
        started_at = timezone.datetime.fromtimestamp(started_at_ms / 1000)
        ended_at = started_at + timezone.timedelta(milliseconds=duration_ms)
    else:
        ended_at = timezone.now()
        started_at = ended_at - timezone.timedelta(milliseconds=duration_ms)

    session_key = getattr(getattr(request, "session", None), "session_key", None) \
                  or request.COOKIES.get("sessionid", "")

    # We want one record per submit, so create() is used
    page_view = PageView.objects.create(
        user=request.user if request.user.is_authenticated else None,
        session_key=session_key or "",
        path=path,
        referrer=data.get("referrer", ""),
        user_agent=request.META.get("HTTP_USER_AGENT", ""),
        started_at=started_at,
        ended_at=ended_at,
        duration_ms=duration_ms,
        extra={
            "visible_only": bool(data.get("visible_only", True)),
            "idle_threshold_ms": data.get("idle_threshold_ms", 120000),
            "trigger": "submit",
        },
    )
    print ("Created page view!")

    return JsonResponse({"ok": True})

@csrf_exempt
# TODO: Fix error where the session key stays the same if a user tries multiple different accounts in the same browser session
# This is causing us to mix up website session data between different users
def track_website_time(request): 
    print ("Successfully got to website time view")
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "POST only"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
        print ("Data received: " + str(data))
    except Exception:
        return JsonResponse({"ok": False, "error": "invalid json"}, status=400)

    session_key = data.get("session_key") or request.COOKIES.get("sessionid", "")
    duration_ms = int(max(0, data.get("duration_ms", 0)))
    print ("About to increment: " + str(duration_ms))

    # Create or update the session
    session, created = WebsiteSession.objects.get_or_create(
        session_key=session_key,
        defaults={
            "user": request.user if request.user.is_authenticated else None,
        },
    )
    print (session)
    print (WebsiteSession.objects.filter(session_key=session_key).count())
    if duration_ms > 0:
        WebsiteSession.objects.filter(pk=session.pk).update(
            total_active_ms=F("total_active_ms") + duration_ms,
            last_ping=timezone.now()
        )

    return JsonResponse({"ok": True})

def track_cursor(request):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "POST only"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
    except json.JSONDecodeError:
        return JsonResponse({"ok": False, "error": "invalid json"}, status=400)

    points = data.get("points", [])
    path = data.get("path", "")
    session_id = data.get("session_id", "")

    user = request.user if request.user.is_authenticated else None

    events = []
    for p in points:
        try:
            events.append(CursorEvent(
                user=user,
                session_id=session_id,
                path=path,
                x=p["x"],
                y=p["y"],
                viewport_width=p["vw"],
                viewport_height=p["vh"],
                scroll_y=p["sy"],
            ))
        except KeyError:
            continue

    if events:
        CursorEvent.objects.bulk_create(events)

    return JsonResponse({"status": "ok", "saved": len(events), "timestamp": timezone.now().isoformat()})

import csv
from django.http import HttpResponse
from openpyxl import Workbook
from .models import CursorEvent

def analyticsCSV(request): # Export analytics data as an XLSX file

    wb = Workbook()

    # Sheet 1: CursorEvent
    sheet = wb.active
    sheet.title = "CursorEvents"

    cursor_qs = CursorEvent.objects.filter(user=request.user).order_by("created_at")

    cursor_headers = [
        "id",
        "created_at",
        "path",
        "session_id",
        "x",
        "y",
        "viewport_width",
        "viewport_height",
        "scroll_y",
    ]
    sheet.append(cursor_headers)

    for e in cursor_qs:
        sheet.append([
            e.id,
            e.created_at.isoformat(),
            e.path,
            e.session_id,
            e.x,
            e.y,
            e.viewport_width,
            e.viewport_height,
            e.scroll_y,
        ])

    # Sheet 2: PageView
    sheet = wb.create_sheet("PageViews")

    page_qs = PageView.objects.filter(user=request.user).order_by("created_at")

    page_headers = [
        "id",
        "session_key",
        "path",
        "referrer",
        "user_agent",
        "started_at",
        "ended_at",
        "duration_ms",
        "extra",
        "created_at",
    ]
    sheet.append(page_headers)

    for p in page_qs:
        sheet.append([
            p.id,
            p.session_key,
            p.path,
            p.referrer,
            p.user_agent,
            p.started_at.isoformat() if p.started_at else "",
            p.ended_at.isoformat() if p.ended_at else "",
            p.duration_ms,
            json.dumps(p.extra),
            p.created_at.isoformat(),
        ])

    # Sheet 3: WebsiteSession
    sheet = wb.create_sheet("WebsiteSessions")

    session_qs = WebsiteSession.objects.filter(user=request.user).order_by("created_at")
    print ("Website session QS length: " + str(session_qs.count()))

    session_headers = [
        "id",
        "session_key",
        "user_id",
        "total_active_ms",
        "last_ping",
        "created_at",
        "updated_at",
    ]
    sheet.append(session_headers)

    for s in session_qs:
        sheet.append([
            s.id,
            s.session_key,
            s.user_id,
            s.total_active_ms,
            s.last_ping.isoformat() if s.last_ping else "",
            s.created_at.isoformat(),
            s.updated_at.isoformat(),
        ])

    # Sheet 4: Deletion Events
    sheet = wb.create_sheet("DeletionEvents")

    deletion_qs = DeletionEvent.objects.filter(user=request.user).order_by("timestamp")

    deletion_headers = [
        "id",
        "session_id",
        "position",
        "length",
        "content",
        "timestamp",
    ]
    sheet.append(deletion_headers)

    for d in deletion_qs:
        sheet.append([
            d.id,
            d.session_id,
            d.position,
            d.length,
            d.content,
            d.timestamp.isoformat(),
        ])

    # Create HTTP that downloads the XLSX file
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="analytics.xlsx"'

    wb.save(response)
    return response

def track_deletion(request):
    print ("Successfully got to deletion view")

    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "POST only"}, status=405)

    try:
        data = json.loads(request.body.decode("utf-8"))
        session_id = int(data.get("session_id", ""))
        user = request.user if request.user.is_authenticated else None
        timestamp = timezone.now()

        DeletionEvent.objects.create(
            user = user,
            session = Session.objects.get(pk=session_id),
            position = data.get("position", 0),
            length = data.get("length", 0),
            content = data.get("content", ""),
            timestamp = timestamp
        )

    except Exception:
        return JsonResponse({"ok": False, "error": "invalid json"}, status=400)
    
    return JsonResponse({"status": "ok", "data_received": data})


